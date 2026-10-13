from typing import List, Dict, Any, Tuple, Optional
from datetime import datetime, date, time
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string
from backend.app.services.formula_parser import format_cell_coordinate


def _color_to_hex(color_obj) -> str:
    """Helper to convert openpyxl color to hex string if available."""
    if not color_obj:
        return ""
    if hasattr(color_obj, 'rgb') and color_obj.rgb:
        rgb_str = str(color_obj.rgb)
        if len(rgb_str) == 8: # AARRGGBB
            return f"#{rgb_str[2:]}"
        elif len(rgb_str) == 6:
            return f"#{rgb_str}"
    return ""

def _format_cell_display(val: Any, num_fmt: Optional[str] = None) -> str:
    """Format cell value for Luckysheet display text."""
    if val is None:
        return ""
    if isinstance(val, (datetime, date)):
        if isinstance(val, datetime):
            return val.strftime("%Y-%m-%d %H:%M:%S") if (val.hour or val.minute or val.second) else val.strftime("%Y-%m-%d")
        return val.strftime("%Y-%m-%d")
    if isinstance(val, time):
        return val.strftime("%H:%M:%S")
    if isinstance(val, (int, float)):
        is_pct = num_fmt and "%" in num_fmt
        is_curr = num_fmt and ("$" in num_fmt or "€" in num_fmt or "£" in num_fmt or "¥" in num_fmt)

        if is_pct:
            return f"{val * 100:.1f}%"
        elif is_curr:
            if isinstance(val, int) or val.is_integer():
                return f"${int(val):,}"
            return f"${val:,.2f}"
        else:
            if isinstance(val, int) or val.is_integer():
                return f"{int(val):,}"
            return f"{val:,.2f}"

    return str(val)

def workbook_to_luckysheet(filepath: str) -> Tuple[List[Dict[str, Any]], int]:
    """
    Converts an Excel workbook into Luckysheet JSON data structure.
    Single-pass extraction using cached values directly from Excel.
    Ensures all cells, formulas, dates, and styles are cleanly JSON-serializable.
    """
    wb_formulas = openpyxl.load_workbook(filepath, data_only=False)
    try:
        wb_values = openpyxl.load_workbook(filepath, data_only=True)
    except Exception:
        wb_values = None

    sheets_data: List[Dict[str, Any]] = []
    total_formulas = 0

    for idx, sheet_name in enumerate(wb_formulas.sheetnames):
        ws_formula = wb_formulas[sheet_name]
        ws_value = wb_values[sheet_name] if wb_values and sheet_name in wb_values.sheetnames else None

        # Handle Chartsheet or non-standard sheets without cell grid
        if not hasattr(ws_formula, "iter_rows") or not hasattr(ws_formula, "max_row"):
            sheet_state = getattr(ws_formula, 'sheet_state', 'visible') or 'visible'
            sheet_obj = {
                "name": sheet_name,
                "index": str(idx),
                "color": "",
                "status": 1 if idx == 0 else 0,
                "order": idx,
                "hide": 1 if sheet_state in ("hidden", "veryHidden") else 0,
                "state": sheet_state,
                "row": 30,
                "column": 15,
                "celldata": [{
                    "r": 0,
                    "c": 0,
                    "v": {
                        "v": f"[{type(ws_formula).__name__}: {sheet_name}]",
                        "m": f"[{type(ws_formula).__name__}: {sheet_name}]",
                        "ct": {"fa": "General", "t": "g"},
                        "bl": 1
                    }
                }],
                "config": {
                    "merge": {},
                    "rowlen": {},
                    "columnlen": {},
                    "rowhidden": {},
                    "colhidden": {}
                }
            }
            sheets_data.append(sheet_obj)
            continue

        celldata: List[Dict[str, Any]] = []
        merge_config: Dict[str, Any] = {}
        row_len_config: Dict[str, int] = {}
        col_len_config: Dict[str, int] = {}
        row_hidden_config: Dict[str, int] = {}
        col_hidden_config: Dict[str, int] = {}

        sheet_max_row = max(ws_formula.max_row or 60, 60)
        sheet_max_col = max(ws_formula.max_column or 26, 26)

        # Extract hidden rows and custom row heights
        for r_num, row_dim in ws_formula.row_dimensions.items():
            if r_num > sheet_max_row + 20:
                continue
            r_idx = r_num - 1
            if r_idx < 0:
                continue
            if row_dim.hidden:
                row_hidden_config[str(r_idx)] = 0
            elif row_dim.height is not None and row_dim.height > 0:
                row_len_config[str(r_idx)] = int(row_dim.height * 1.33)

        # Extract hidden columns and custom column widths
        for col_key, col_dim in ws_formula.column_dimensions.items():
            min_c = col_dim.min
            max_c = col_dim.max
            if min_c is None:
                try:
                    min_c = column_index_from_string(str(col_key))
                except Exception:
                    min_c = None
            if max_c is None:
                max_c = min_c

            if min_c is not None and max_c is not None:
                effective_max_c = min(max_c, sheet_max_col + 10)
                if min_c > effective_max_c:
                    continue
                for c_num in range(min_c, effective_max_c + 1):
                    c_idx = c_num - 1
                    if c_idx < 0:
                        continue
                    if col_dim.hidden or (col_dim.width is not None and col_dim.width == 0):
                        col_hidden_config[str(c_idx)] = 0
                    elif col_dim.width is not None and col_dim.width > 0:
                        col_len_config[str(c_idx)] = max(20, int(col_dim.width * 8 + 5))


        # Merge cells
        for merge_range in ws_formula.merged_cells.ranges:
            min_col, min_row, max_col, max_row = merge_range.bounds
            r_idx = min_row - 1
            c_idx = min_col - 1
            row_span = (max_row - min_row) + 1
            col_span = (max_col - min_col) + 1
            merge_config[f"{r_idx}_{c_idx}"] = {
                "r": r_idx,
                "c": c_idx,
                "rs": row_span,
                "cs": col_span
            }


        # Iterate over populated cells in single pass
        for row in ws_formula.iter_rows():
            for cell_f in row:
                if cell_f.value is None and cell_f.comment is None:
                    continue

                r_idx = cell_f.row - 1
                c_idx = cell_f.column - 1
                coord = format_cell_coordinate(cell_f.row, cell_f.column)

                val_f = cell_f.value
                val_cached = None
                if ws_value:
                    try:
                        val_cached = ws_value.cell(row=cell_f.row, column=cell_f.column).value
                    except Exception:
                        val_cached = None

                # Extract formula if present (handling string formulas and openpyxl ArrayFormula)
                is_formula = False
                formula_str = None
                if isinstance(val_f, str) and val_f.startswith("="):
                    is_formula = True
                    formula_str = val_f
                elif hasattr(val_f, "text"):
                    # openpyxl ArrayFormula object
                    f_text = str(val_f.text)
                    if not f_text.startswith("="):
                        f_text = "=" + f_text
                    is_formula = True
                    formula_str = f_text

                if is_formula:
                    total_formulas += 1

                # Unwrap val_cached if it is an object with .text
                if hasattr(val_cached, "text"):
                    val_cached = str(val_cached.text)

                # Determine effective calculated value (prefer Excel cached value)
                if is_formula:
                    if val_cached is not None and not (isinstance(val_cached, str) and val_cached.startswith("=")):
                        final_value = val_cached
                    else:
                        final_value = formula_str
                else:
                    final_value = val_f

                # Sanitize final_value for JSON serialization & Luckysheet
                cell_type = "g"
                if isinstance(final_value, (int, float)):
                    cell_type = "n"
                elif isinstance(final_value, (datetime, date)):
                    if isinstance(final_value, datetime):
                        final_value = final_value.strftime("%Y-%m-%d %H:%M:%S") if (final_value.hour or final_value.minute or final_value.second) else final_value.strftime("%Y-%m-%d")
                    else:
                        final_value = final_value.strftime("%Y-%m-%d")
                    cell_type = "d"
                elif isinstance(final_value, time):
                    final_value = final_value.strftime("%H:%M:%S")
                    cell_type = "d"
                elif isinstance(final_value, str):
                    cell_type = "s"
                elif final_value is not None:
                    final_value = str(final_value)

                display_text = _format_cell_display(final_value, cell_f.number_format)

                cell_obj: Dict[str, Any] = {
                    "v": final_value,
                    "m": display_text,
                    "ct": {
                        "fa": cell_f.number_format or "General",
                        "t": cell_type
                    }
                }

                if formula_str:
                    cell_obj["f"] = formula_str

                # Styles
                if cell_f.font:
                    if cell_f.font.bold:
                        cell_obj["bl"] = 1
                    if cell_f.font.italic:
                        cell_obj["it"] = 1
                    if cell_f.font.size:
                        cell_obj["fs"] = int(cell_f.font.size)
                    fc = _color_to_hex(cell_f.font.color)
                    if fc:
                        cell_obj["fc"] = fc

                if cell_f.fill and hasattr(cell_f.fill, 'start_color'):
                    bg = _color_to_hex(cell_f.fill.start_color)
                    if bg and bg != "#000000":
                        cell_obj["bg"] = bg

                celldata.append({
                    "r": r_idx,
                    "c": c_idx,
                    "v": cell_obj
                })

        sheet_state = getattr(ws_formula, 'sheet_state', 'visible') or 'visible'
        sheet_obj = {
            "name": sheet_name,
            "index": str(idx),
            "color": "",
            "status": 1 if idx == 0 else 0,
            "order": idx,
            "hide": 1 if sheet_state in ("hidden", "veryHidden") else 0,
            "state": sheet_state,
            "row": max(ws_formula.max_row or 60, 60),
            "column": max(ws_formula.max_column or 26, 26),
            "celldata": celldata,
            "config": {
                "merge": merge_config,
                "rowlen": row_len_config,
                "columnlen": col_len_config,
                "rowhidden": row_hidden_config,
                "colhidden": col_hidden_config
            }
        }
        sheets_data.append(sheet_obj)

    wb_formulas.close()
    if wb_values:
        wb_values.close()

    return sheets_data, total_formulas
